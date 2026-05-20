"""Excel training report: input sanity and backward math verification."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from sklearn.metrics import silhouette_score

from telekom_profiler.ml.features import build_subscriber_features, feature_matrix, load_usage_panel
from telekom_profiler.ml.schema import (
    CLUSTER_FEATURES,
    MONTH_COL,
    RAW_NUMERIC_COLS,
    SUBSCRIBER_ID_COL,
    TREND_COLS,
)


def _audit_row(
    check: str,
    *,
    category: str,
    formula: str,
    status: str,
    unit_note: str = "",
    reverse_check: str = "",
    bound_expected: str = "",
    verdict: str = "",
    max_abs_error: float | None = None,
    n_failures: int | None = None,
    n_total: int | None = None,
    sample_subscriber: str = "",
    sample_inputs: str = "",
    sample_computed: str = "",
    sample_saved: str = "",
    sample_reverse: str = "",
    details: str = "",
) -> dict[str, object]:
    if not verdict:
        verdict = "VALID" if status == "PASS" else "ERROR"
    elif status == "FAIL" and verdict == "VALID":
        verdict = "ERROR"
    return {
        "category": category,
        "check": check,
        "formula": formula,
        "unit_note": unit_note,
        "reverse_check": reverse_check,
        "bound_expected": bound_expected,
        "status": status,
        "verdict": verdict,
        "max_abs_error": max_abs_error,
        "n_failures": n_failures,
        "n_total": n_total,
        "sample_subscriber": sample_subscriber,
        "sample_inputs": sample_inputs,
        "sample_computed": sample_computed,
        "sample_saved": sample_saved,
        "sample_reverse": sample_reverse,
        "details": details,
    }


def _fmt_num(value: float) -> str:
    if abs(value) >= 1000 or (0 < abs(value) < 0.001):
        return f"{value:.6g}"
    return f"{value:.4f}"


def _sample_triplet(
    features: pd.DataFrame,
    sid: str,
    *,
    inputs: str,
    computed: float,
    saved_col: str,
    reverse: str,
) -> dict[str, str]:
    saved = float(features.loc[features[SUBSCRIBER_ID_COL] == sid, saved_col].iloc[0])
    return {
        "sample_subscriber": sid,
        "sample_inputs": inputs,
        "sample_computed": _fmt_num(computed),
        "sample_saved": _fmt_num(saved),
        "sample_reverse": reverse,
    }


def _bounds_note(series: pd.Series, lo: float | None = None, hi: float | None = None) -> str:
    parts = [f"observed [{series.min():.4g}, {series.max():.4g}]"]
    if lo is not None or hi is not None:
        lo_s = f"{lo}" if lo is not None else "—"
        hi_s = f"{hi}" if hi is not None else "—"
        parts.append(f"expected [{lo_s}, {hi_s}]")
    return "; ".join(parts)


def build_input_sanity_rows(panel: pd.DataFrame) -> list[dict[str, object]]:
    """Schema and range checks on the raw monthly panel."""
    rows: list[dict[str, object]] = []
    numeric_cols = [MONTH_COL, *RAW_NUMERIC_COLS]

    extra_text_cols = [c for c in panel.columns if c not in {SUBSCRIBER_ID_COL, *numeric_cols}]
    rows.append(
        _audit_row(
            "No non-numeric training feature columns",
            category="input_schema",
            formula="columns == {subscriber_id, month} ∪ RAW_NUMERIC_COLS",
            unit_note="K-Means inputs must be numeric only",
            reverse_check="N/A (schema)",
            bound_expected="no seed_archetype / text labels",
            status="PASS" if not extra_text_cols else "FAIL",
            details="none" if not extra_text_cols else ", ".join(extra_text_cols),
        )
    )

    missing_numeric = int(panel[numeric_cols].isna().sum().sum())
    rows.append(
        _audit_row(
            "No nulls in numeric columns",
            category="input_schema",
            formula="∀ cell in numeric cols: not NaN",
            status="PASS" if missing_numeric == 0 else "FAIL",
            n_failures=missing_numeric,
            details=f"null_count={missing_numeric}",
        )
    )

    months_ok = bool(panel[MONTH_COL].between(1, 12).all())
    rows.append(
        _audit_row(
            "Month index in [1, 12]",
            category="input_schema",
            formula="1 ≤ month ≤ 12",
            status="PASS" if months_ok else "FAIL",
            details=f"min={int(panel[MONTH_COL].min())}, max={int(panel[MONTH_COL].max())}",
        )
    )

    non_negative_cols = [
        c for c in RAW_NUMERIC_COLS if c not in {"night_usage_ratio", "weekend_usage_ratio"}
    ]
    negatives = int((panel[non_negative_cols] < 0).sum().sum())
    rows.append(
        _audit_row(
            "Non-negative usage values",
            category="input_ranges",
            formula="usage cols ≥ 0",
            status="PASS" if negatives == 0 else "FAIL",
            n_failures=negatives,
            details=f"negative_cells={negatives}",
        )
    )

    ratio_cols = ["night_usage_ratio", "weekend_usage_ratio"]
    ratio_ok = bool(panel[ratio_cols].apply(lambda s: s.between(0, 1).all()).all())
    rows.append(
        _audit_row(
            "Usage ratio columns in [0, 1]",
            category="input_ranges",
            formula="0 ≤ night/weekend ratio ≤ 1",
            status="PASS" if ratio_ok else "FAIL",
            details=", ".join(ratio_cols),
        )
    )

    active_le_total = bool((panel["lines_active"] <= panel["lines_total"]).all())
    rows.append(
        _audit_row(
            "lines_active <= lines_total",
            category="input_ranges",
            formula="lines_active ≤ lines_total (per month)",
            status="PASS" if active_le_total else "FAIL",
            details="validated row-wise",
        )
    )
    return rows


def _compare_series(
    saved: pd.Series,
    rebuilt: pd.Series,
    *,
    rtol: float = 1e-5,
    atol: float = 1e-4,
) -> tuple[int, float, int]:
    diff = (saved.astype(float) - rebuilt.astype(float)).abs()
    mask = ~np.isclose(saved.astype(float), rebuilt.astype(float), rtol=rtol, atol=atol)
    n_fail = int(mask.sum())
    max_err = float(diff.max()) if len(diff) else 0.0
    return n_fail, max_err, len(diff)


def build_math_backward_rows(
    panel: pd.DataFrame,
    features: pd.DataFrame,
    cluster_map: pd.DataFrame,
    artifacts_dir: Path,
    summary: dict[str, object],
    *,
    rtol: float = 1e-5,
    atol: float = 1e-4,
) -> tuple[list[dict[str, object]], pd.DataFrame]:
    """
    Recompute pipeline steps from raw panel and compare to saved artifacts.

    Returns aggregate check rows and a detail frame of per-subscriber/column diffs (top offenders).
    """
    rows: list[dict[str, object]] = []
    detail_parts: list[pd.DataFrame] = []

    n_subscribers = panel[SUBSCRIBER_ID_COL].nunique()
    months_per_sub = panel.groupby(SUBSCRIBER_ID_COL)[MONTH_COL].nunique()
    bad_month_counts = int((months_per_sub != 12).sum())
    rows.append(
        _audit_row(
            "12 months per subscriber",
            category="aggregation",
            formula="count(month) per subscriber_id == 12",
            status="PASS" if bad_month_counts == 0 else "FAIL",
            n_failures=bad_month_counts,
            n_total=n_subscribers,
            details=f"violations={bad_month_counts}",
        )
    )

    rebuilt = build_subscriber_features(panel)
    merged = features.merge(
        rebuilt,
        on=SUBSCRIBER_ID_COL,
        how="inner",
        suffixes=("_saved", "_rebuilt"),
    )
    if len(merged) != len(features):
        rows.append(
            _audit_row(
                "Subscriber id alignment (features rebuild)",
                category="aggregation",
                formula="saved features ⟷ rebuilt features (inner join)",
                status="FAIL",
                n_failures=abs(len(features) - len(merged)),
                n_total=len(features),
                details=f"saved={len(features)}, matched={len(merged)}",
            )
        )

    verify_cols = list(CLUSTER_FEATURES) + list(TREND_COLS)
    total_fail = 0
    global_max_err = 0.0
    for col in verify_cols:
        saved_col = f"{col}_saved"
        rebuilt_col = f"{col}_rebuilt"
        if saved_col not in merged.columns:
            continue
        n_fail, max_err, n_total = _compare_series(merged[saved_col], merged[rebuilt_col], rtol=rtol, atol=atol)
        total_fail += n_fail
        global_max_err = max(global_max_err, max_err)
        if n_fail > 0:
            bad = merged.loc[
                ~np.isclose(merged[saved_col], merged[rebuilt_col], rtol=rtol, atol=atol),
                [SUBSCRIBER_ID_COL, saved_col, rebuilt_col],
            ].copy()
            bad["column"] = col
            bad["abs_error"] = (bad[saved_col] - bad[rebuilt_col]).abs()
            bad = bad.rename(columns={saved_col: "saved", rebuilt_col: "rebuilt"})
            detail_parts.append(bad[[SUBSCRIBER_ID_COL, "column", "saved", "rebuilt", "abs_error"]])

    rows.append(
        _audit_row(
            "Feature matrix matches panel rebuild",
            category="aggregation",
            formula="build_subscriber_features(panel) == subscriber_features.csv",
            status="PASS" if total_fail == 0 else "FAIL",
            max_abs_error=global_max_err,
            n_failures=total_fail,
            n_total=len(merged) * len(verify_cols),
            details=f"checked {len(verify_cols)} columns, rtol={rtol}, atol={atol}",
        )
    )

    # Derived formula spot-checks on saved features (inverse definitions)
    lt = features["lines_total_mean"].clip(lower=1.0)
    la = features["lines_active_mean"].clip(lower=0.1)
    data = features["data_gb_mean"]
    roam = features["roaming_days_mean"]
    countries = features["countries_visited_mean"]
    active_days = features["active_days_mean"].clip(lower=1.0)
    plan = features["plan_tier_mean"]

    expected_active_ratio = la / lt
    n_fail, max_err, _ = _compare_series(
        features["active_line_ratio"], expected_active_ratio, rtol=rtol, atol=atol
    )
    sid = "SUB00002"
    r2 = features.loc[features[SUBSCRIBER_ID_COL] == sid].iloc[0]
    lt2, la2 = max(r2["lines_total_mean"], 1.0), max(r2["lines_active_mean"], 0.1)
    comp_active = la2 / lt2
    rows.append(
        _audit_row(
            "active_line_ratio = lines_active_mean / lines_total_mean",
            category="derived",
            formula="active_line_ratio = la / lt (lt≥1, la≥0.1)",
            unit_note="Ratio 0–1 (not ×100 %)",
            reverse_check="la = active_line_ratio × lt",
            bound_expected=_bounds_note(features["active_line_ratio"], 0, 1),
            status="PASS" if n_fail == 0 else "FAIL",
            max_abs_error=max_err,
            n_failures=n_fail,
            n_total=len(features),
            **_sample_triplet(
                features,
                sid,
                inputs=f"lt={lt2:.4f}, la={la2:.4f}",
                computed=comp_active,
                saved_col="active_line_ratio",
                reverse=f"la_back={comp_active * lt2:.4f}",
            ),
        )
    )

    expected_roam_ratio = roam / 30.0
    n_fail, max_err, _ = _compare_series(
        features["roaming_days_ratio"], expected_roam_ratio, rtol=rtol, atol=atol
    )
    sid3 = "SUB00003"
    r3 = features.loc[features[SUBSCRIBER_ID_COL] == sid3].iloc[0]
    comp_roam = float(r3["roaming_days_mean"]) / 30.0
    rows.append(
        _audit_row(
            "roaming_days_ratio = roaming_days_mean / 30",
            category="derived",
            formula="roaming_days_ratio = roam / 30",
            unit_note="Days per 30-day month (not ×100 %)",
            reverse_check="roam_days = roaming_days_ratio × 30",
            bound_expected=_bounds_note(features["roaming_days_ratio"], 0, 1),
            status="PASS" if n_fail == 0 else "FAIL",
            max_abs_error=max_err,
            n_failures=n_fail,
            **_sample_triplet(
                features,
                sid3,
                inputs=f"roam={r3['roaming_days_mean']:.4f}",
                computed=comp_roam,
                saved_col="roaming_days_ratio",
                reverse=f"roam_back={comp_roam * 30:.4f}",
            ),
        )
    )

    expected_roam_intensity = features["roaming_days_ratio"] * countries
    n_fail, max_err, _ = _compare_series(
        features["roaming_intensity"], expected_roam_intensity, rtol=rtol, atol=atol
    )
    comp_int = comp_roam * float(r3["countries_visited_mean"])
    rows.append(
        _audit_row(
            "roaming_intensity = roaming_days_ratio × countries_visited_mean",
            category="derived",
            formula="roaming_intensity = roaming_days_ratio × countries",
            unit_note="Composite score (unbounded)",
            reverse_check="countries = intensity / roaming_days_ratio",
            bound_expected=_bounds_note(features["roaming_intensity"]),
            status="PASS" if n_fail == 0 else "FAIL",
            max_abs_error=max_err,
            n_failures=n_fail,
            **_sample_triplet(
                features,
                sid3,
                inputs=f"ratio={comp_roam:.4f}, countries={r3['countries_visited_mean']:.4f}",
                computed=comp_int,
                saved_col="roaming_intensity",
                reverse=f"countries_back={comp_int / comp_roam:.4f}" if comp_roam else "N/A",
            ),
        )
    )

    expected_pct_idle = ((lt - la) / lt).clip(0.0, 1.0)
    n_fail, max_err, _ = _compare_series(
        features["pct_idle_lines"], expected_pct_idle, rtol=rtol, atol=atol
    )
    comp_idle = float(np.clip((lt2 - la2) / lt2, 0.0, 1.0))
    rows.append(
        _audit_row(
            "pct_idle_lines = (lines_total − lines_active) / lines_total",
            category="derived",
            formula="pct_idle_lines = clip((lt−la)/lt, 0, 1)",
            unit_note="Fraction 0–1 idle lines (display ×100 for %)",
            reverse_check="(lt−la) = pct_idle × lt",
            bound_expected=_bounds_note(features["pct_idle_lines"], 0, 1),
            status="PASS" if n_fail == 0 else "FAIL",
            max_abs_error=max_err,
            n_failures=n_fail,
            **_sample_triplet(
                features,
                sid,
                inputs=f"lt={lt2:.4f}, la={la2:.4f}",
                computed=comp_idle,
                saved_col="pct_idle_lines",
                reverse=f"idle_lines_back={comp_idle * lt2:.4f} vs lt−la={lt2 - la2:.4f}",
            ),
        )
    )

    expected_session = features["avg_session_mb_mean"] * data
    n_fail, max_err, _ = _compare_series(
        features["session_intensity"], expected_session, rtol=rtol, atol=atol
    )
    comp_sess = float(r2["avg_session_mb_mean"]) * float(r2["data_gb_mean"])
    rows.append(
        _audit_row(
            "session_intensity = avg_session_mb_mean × data_gb_mean",
            category="derived",
            formula="session_intensity = avg_session_mb × data_gb",
            unit_note="MB × GB composite (unbounded)",
            reverse_check="avg_session_mb = intensity / data_gb",
            bound_expected=_bounds_note(features["session_intensity"]),
            status="PASS" if n_fail == 0 else "FAIL",
            max_abs_error=max_err,
            n_failures=n_fail,
            **_sample_triplet(
                features,
                sid,
                inputs=f"session_mb={r2['avg_session_mb_mean']:.2f}, data={r2['data_gb_mean']:.2f}",
                computed=comp_sess,
                saved_col="session_intensity",
                reverse=f"session_mb_back={comp_sess / r2['data_gb_mean']:.2f}",
            ),
        )
    )

    expected_plan_gap = (plan - (data / 50.0)).clip(lower=0.0)
    n_fail, max_err, _ = _compare_series(
        features["plan_usage_gap"], expected_plan_gap, rtol=rtol, atol=atol
    )
    comp_gap = max(0.0, float(r2["plan_tier_mean"]) - float(r2["data_gb_mean"]) / 50.0)
    rows.append(
        _audit_row(
            "plan_usage_gap = max(0, plan_tier_mean − data_gb_mean/50)",
            category="derived",
            formula="plan_usage_gap = max(0, plan − data/50)",
            unit_note="Tier minus usage index (≥0)",
            reverse_check="plan = gap + data/50 when gap>0",
            bound_expected=_bounds_note(features["plan_usage_gap"], 0, None),
            status="PASS" if n_fail == 0 else "FAIL",
            max_abs_error=max_err,
            n_failures=n_fail,
            **_sample_triplet(
                features,
                sid,
                inputs=f"plan={r2['plan_tier_mean']:.4f}, data={r2['data_gb_mean']:.2f}",
                computed=comp_gap,
                saved_col="plan_usage_gap",
                reverse=f"plan_back={comp_gap + r2['data_gb_mean'] / 50:.4f}",
            ),
        )
    )

    active = features["active_line_ratio"]
    for name, formula, unit, reverse, expected_series in [
        (
            "data_per_active_line = data_gb_mean / active_line_ratio",
            "data / active_line_ratio",
            "GB per active line",
            "data = data_per_active_line × active_line_ratio",
            features["data_gb_mean"] / active,
        ),
        (
            "sms_per_gb = sms_count_mean / max(data_gb_mean, 0.5)",
            "sms / max(data, 0.5)",
            "SMS per GB",
            "sms = sms_per_gb × data",
            features["sms_count_mean"] / features["data_gb_mean"].clip(lower=0.5),
        ),
        (
            "voice_per_active_day = voice_min_mean / active_days_mean",
            "voice / active_days",
            "min per active day",
            "voice = voice_per_active_day × active_days",
            features["voice_min_mean"] / features["active_days_mean"].clip(lower=1.0),
        ),
        (
            "data_per_active_day = data_gb_mean / active_days_mean",
            "data / active_days",
            "GB per active day",
            "data = data_per_active_day × active_days",
            features["data_gb_mean"] / features["active_days_mean"].clip(lower=1.0),
        ),
    ]:
        col = name.split(" = ")[0]
        n_fail, max_err, _ = _compare_series(features[col], expected_series, rtol=rtol, atol=atol)
        rows.append(
            _audit_row(
                name,
                category="derived",
                formula=formula,
                unit_note=unit,
                reverse_check=reverse,
                bound_expected=_bounds_note(features[col]),
                status="PASS" if n_fail == 0 else "FAIL",
                max_abs_error=max_err,
                n_failures=n_fail,
                n_total=len(features),
            )
        )

    # K-Means / scaler backward checks
    kmeans = joblib.load(artifacts_dir / "kmeans.pkl")
    scaler = joblib.load(artifacts_dir / "scaler.pkl")
    label_map_raw = json.loads((artifacts_dir / "label_map.json").read_text(encoding="utf-8"))
    label_map = {int(k): v for k, v in label_map_raw.items()}
    cluster_features = tuple(
        json.loads((artifacts_dir / "cluster_features.json").read_text(encoding="utf-8"))
    )

    x, subscriber_ids = feature_matrix(features, columns=cluster_features)
    x_scaled = scaler.transform(x)
    centroids_scaled = kmeans.cluster_centers_
    labels = kmeans.predict(x_scaled)

    dist_fail = 0
    idx_fail = 0
    label_fail = 0
    primary_dist_fail = 0
    dist_max_err = 0.0

    cm = cluster_map.set_index(SUBSCRIBER_ID_COL)
    for i, sid in enumerate(subscriber_ids):
        point = x_scaled[i : i + 1]
        dists = np.linalg.norm(centroids_scaled - point, axis=1)
        order = np.argsort(dists)
        primary_idx = int(order[0])
        primary_dist = float(dists[primary_idx])

        saved_idx = int(cm.loc[sid, "cluster_idx"])
        if saved_idx != primary_idx or int(labels[i]) != primary_idx:
            idx_fail += 1
        if not np.isclose(float(cm.loc[sid, "primary_distance"]), primary_dist, rtol=rtol, atol=atol):
            primary_dist_fail += 1
            dist_max_err = max(dist_max_err, abs(float(cm.loc[sid, "primary_distance"]) - primary_dist))
        if cm.loc[sid, "cluster_label"] != label_map[primary_idx]:
            label_fail += 1

        for j in range(len(label_map)):
            col_name = f"dist_{label_map[j]}"
            if col_name in cm.columns:
                saved_d = float(cm.loc[sid, col_name])
                if not np.isclose(saved_d, float(dists[j]), rtol=rtol, atol=atol):
                    dist_fail += 1
                    dist_max_err = max(dist_max_err, abs(saved_d - float(dists[j])))

    n = len(subscriber_ids)
    rows.append(
        _audit_row(
            "cluster_idx = argmin_euclidean(scaled features, centroids)",
            category="clustering",
            formula="cluster_idx = argmin ||x_scaled - centroid_k||₂",
            status="PASS" if idx_fail == 0 else "FAIL",
            n_failures=idx_fail,
            n_total=n,
        )
    )
    rows.append(
        _audit_row(
            "primary_distance matches recomputed min distance",
            category="clustering",
            formula="primary_distance = min_k ||x_scaled - centroid_k||₂",
            status="PASS" if primary_dist_fail == 0 else "FAIL",
            max_abs_error=dist_max_err,
            n_failures=primary_dist_fail,
            n_total=n,
        )
    )
    rows.append(
        _audit_row(
            "dist_* columns match all centroid distances",
            category="clustering",
            formula="dist_label_k = ||x_scaled - centroid_k||₂",
            status="PASS" if dist_fail == 0 else "FAIL",
            max_abs_error=dist_max_err,
            n_failures=dist_fail,
            n_total=n * len(label_map),
        )
    )
    rows.append(
        _audit_row(
            "cluster_label = label_map[cluster_idx]",
            category="clustering",
            formula="cluster_label = label_map[cluster_idx]",
            status="PASS" if label_fail == 0 else "FAIL",
            n_failures=label_fail,
            n_total=n,
        )
    )

    sil_recomputed = float(silhouette_score(x_scaled, labels))
    sil_saved = float(summary["silhouette"])
    sil_ok = bool(np.isclose(sil_recomputed, sil_saved, rtol=1e-4, atol=1e-4))
    rows.append(
        _audit_row(
            "Silhouette score matches recomputation",
            category="clustering",
            formula="silhouette(x_scaled, kmeans.labels) == training summary",
            status="PASS" if sil_ok else "FAIL",
            max_abs_error=abs(sil_recomputed - sil_saved),
            details=f"saved={sil_saved:.6f}, recomputed={sil_recomputed:.6f}",
        )
    )

    # Confidence backward check (same thresholds as train._confidence)
    conf_fail = 0
    for _, r in cluster_map.iterrows():
        p = float(r["primary_distance"])
        s = float(r["secondary_distance"])
        if s < 1e-9:
            expected = "High"
        else:
            ratio = p / s
            if ratio < 0.75:
                expected = "High"
            elif ratio < 0.9:
                expected = "Medium"
            else:
                expected = "Low"
        if r["confidence"] != expected:
            conf_fail += 1
    conf_dist = cluster_map["confidence"].value_counts()
    conf_verdict = "VALID" if conf_fail == 0 else "ERROR"
    conf_details = f"High={conf_dist.get('High', 0)}, Medium={conf_dist.get('Medium', 0)}, Low={conf_dist.get('Low', 0)}"
    if conf_fail == 0 and len(conf_dist) == 1 and conf_dist.index[0] == "High":
        conf_verdict = "WARNING"
        conf_details += "; all subscribers High — labels not discriminative on this dataset"
    r1 = cm.loc["SUB00001"]
    p1, s1 = float(r1["primary_distance"]), float(r1["secondary_distance"])
    ratio1 = p1 / s1 if s1 > 1e-9 else 0.0
    rows.append(
        _audit_row(
            "confidence label from primary/secondary distance ratio",
            category="clustering",
            formula="ratio = primary_dist / secondary_dist; High if <0.75, Medium if <0.9",
            unit_note="Distance ratio (not ×100 %)",
            reverse_check="secondary_dist = primary_dist / ratio",
            bound_expected="ratio typically 0–1 on separated clusters",
            status="PASS" if conf_fail == 0 else "FAIL",
            verdict=conf_verdict,
            n_failures=conf_fail,
            n_total=len(cluster_map),
            sample_subscriber="SUB00001",
            sample_inputs=f"p={p1:.4f}, s={s1:.4f}",
            sample_computed=_fmt_num(ratio1),
            sample_saved=str(r1["confidence"]),
            sample_reverse=f"s_back={p1 / ratio1:.4f}" if ratio1 else "N/A",
            details=conf_details,
        )
    )

    detail_df = pd.DataFrame()
    if detail_parts:
        detail_df = pd.concat(detail_parts, ignore_index=True)
        detail_df = detail_df.sort_values("abs_error", ascending=False).head(100)

    return rows, detail_df


_AUDIT_COLUMN_ORDER = [
    "category",
    "check",
    "formula",
    "unit_note",
    "reverse_check",
    "bound_expected",
    "status",
    "verdict",
    "max_abs_error",
    "n_failures",
    "n_total",
    "sample_subscriber",
    "sample_inputs",
    "sample_computed",
    "sample_saved",
    "sample_reverse",
    "details",
]


def _order_audit_df(df: pd.DataFrame) -> pd.DataFrame:
    cols = [c for c in _AUDIT_COLUMN_ORDER if c in df.columns]
    extra = [c for c in df.columns if c not in cols]
    return df[cols + extra]


def _row_style(status: str | None, verdict: str | None) -> tuple[PatternFill | None, Font | None]:
    """Row fill and bold font for status/verdict cells: green PASS/VALID, yellow WARNING, red FAIL/ERROR."""
    from openpyxl.styles import Font, PatternFill

    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    warn_fill = PatternFill("solid", fgColor="FFEB9C")

    label = (verdict or "").strip().upper() or (status or "").strip().upper()
    if label in ("VALID", "PASS"):
        return pass_fill, Font(bold=True, color="006100")
    if label == "WARNING":
        return warn_fill, Font(bold=True, color="9C6500")
    if label in ("ERROR", "FAIL"):
        return fail_fill, Font(bold=True, color="9C0006")
    return None, None


def _format_excel_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="E20074")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    styled_sheets = {"sanity_input", "sanity_math", "sanity_math_detail"}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 1:
            continue

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        headers = {cell.value: cell.column for cell in ws[1] if cell.value}
        status_col = headers.get("status")
        verdict_col = headers.get("verdict")

        for row_idx in range(2, ws.max_row + 1):
            st = ws.cell(row=row_idx, column=status_col).value if status_col else None
            vd = ws.cell(row=row_idx, column=verdict_col).value if verdict_col else None
            row_fill, label_font = _row_style(
                str(st) if st is not None else None,
                str(vd) if vd is not None else None,
            )

            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.alignment = wrap
                cell.border = border
                if sheet_name in styled_sheets and row_fill is not None:
                    cell.fill = row_fill

            if label_font and sheet_name in styled_sheets:
                if status_col:
                    c = ws.cell(row=row_idx, column=status_col)
                    c.font = label_font
                if verdict_col:
                    c = ws.cell(row=row_idx, column=verdict_col)
                    c.font = label_font

        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            header = ws.cell(row=1, column=col_idx).value
            max_len = len(str(header or ""))
            for row_idx in range(2, min(ws.max_row, 80) + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, min(len(str(val)), 48))
            ws.column_dimensions[letter].width = max(10, min(max_len + 2, 42))

        ws.freeze_panes = "A2"
        if sheet_name == "sanity_math":
            ws.auto_filter.ref = ws.dimensions

    wb.save(path)


def write_training_excel_report(
    input_csv: Path,
    artifacts_dir: Path,
    summary: dict[str, object],
    out_xlsx: Path,
) -> None:
    """Write multi-sheet Excel report with input and backward math sanity checks."""
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for Excel reports. Install ML dependencies:\n"
            "  pip install -r requirements-ml.txt\n"
            "  # or: pip install -e \".[ml]\""
        ) from exc

    panel = load_usage_panel(input_csv)
    features = pd.read_csv(artifacts_dir / "subscriber_features.csv")
    cluster_map = pd.read_csv(artifacts_dir / "subscriber_cluster_map.csv")

    input_sanity = _order_audit_df(pd.DataFrame(build_input_sanity_rows(panel)))
    math_rows, math_detail = build_math_backward_rows(
        panel, features, cluster_map, artifacts_dir, summary
    )
    math_sanity = _order_audit_df(pd.DataFrame(math_rows))

    overview_df = pd.DataFrame(
        [
            {"metric": "input_csv", "value": str(input_csv)},
            {"metric": "n_rows_monthly_panel", "value": len(panel)},
            {"metric": "n_subscribers", "value": int(summary["n_subscribers"])},
            {"metric": "silhouette", "value": float(summary["silhouette"])},
            {"metric": "math_checks_passed", "value": int((math_sanity["status"] == "PASS").sum())},
            {"metric": "math_checks_failed", "value": int((math_sanity["status"] == "FAIL").sum())},
            {"metric": "clusters", "value": cluster_map["cluster_label"].nunique()},
            {"metric": "label_map", "value": json.dumps(summary.get("label_map", {}), ensure_ascii=True)},
        ]
    )
    cluster_counts = (
        cluster_map["cluster_label"].value_counts().rename_axis("cluster_label").reset_index(name="count")
    )
    confidence_counts = (
        cluster_map["confidence"].value_counts().rename_axis("confidence").reset_index(name="count")
    )
    numeric_summary = features.describe().T.reset_index().rename(columns={"index": "feature"})

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        overview_df.to_excel(writer, sheet_name="overview", index=False)
        input_sanity.to_excel(writer, sheet_name="sanity_input", index=False)
        math_sanity.to_excel(writer, sheet_name="sanity_math", index=False)
        if not math_detail.empty:
            math_detail.to_excel(writer, sheet_name="sanity_math_detail", index=False)
        cluster_counts.to_excel(writer, sheet_name="cluster_counts", index=False)
        confidence_counts.to_excel(writer, sheet_name="confidence_counts", index=False)
        numeric_summary.to_excel(writer, sheet_name="feature_summary", index=False)

    _format_excel_workbook(out_xlsx)
